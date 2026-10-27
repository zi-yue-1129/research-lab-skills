"""Tests for the lecture-slides deck builder."""

from __future__ import annotations

import importlib.util
from pathlib import Path

import pytest
import yaml
from pptx import Presentation

_SKILL_DIR = Path(__file__).resolve().parents[2]
_spec = importlib.util.spec_from_file_location(
    "lecture_build_deck", _SKILL_DIR / "scripts" / "build_deck.py")
build_deck = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(build_deck)

GENERIC_DEMO = _SKILL_DIR / "examples" / "generic-demo.yaml"
DEMO = _SKILL_DIR / "presets" / "zh-classroom" / "examples" / "knn-demo.yaml"


def _write(tmp_path: Path, spec: dict) -> Path:
    path = tmp_path / "deck.yaml"
    path.write_text(yaml.safe_dump(spec, allow_unicode=True), encoding="utf-8")
    return path


def _texts(slide) -> list[str]:
    return [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]


@pytest.mark.parametrize("theme", ["warm-paper", "navy-academic"])
def test_demo_builds_every_layout_in_both_themes(tmp_path: Path, theme: str) -> None:
    out = tmp_path / "demo.pptx"
    warnings = build_deck.build(DEMO, out, theme)
    prs = Presentation(str(out))
    spec = yaml.safe_load(DEMO.read_text(encoding="utf-8"))
    assert len(prs.slides) == len(spec["slides"])
    assert {s["layout"] for s in spec["slides"]} == set(build_deck.LAYOUTS)
    # The demo ships without images, so placeholders are the only warnings.
    assert all("image placeholder" in w for w in warnings), warnings


def test_every_demo_slide_carries_notes_and_page_number(tmp_path: Path) -> None:
    out = tmp_path / "demo.pptx"
    build_deck.build(DEMO, out)
    prs = Presentation(str(out))
    total = len(prs.slides)
    for i, slide in enumerate(prs.slides, 1):
        assert slide.notes_slide.notes_text_frame.text.strip()
        if i > 1:
            assert f"{i} / {total}" in _texts(slide)


def test_theme_background_and_fonts_are_applied(tmp_path: Path) -> None:
    out = tmp_path / "demo.pptx"
    build_deck.build(DEMO, out, "warm-paper")
    slide = Presentation(str(out)).slides[1]
    assert str(slide.background.fill.fore_color.rgb) == "FFF8EC"
    xml = slide.shapes._spTree.xml
    assert 'typeface="Times New Roman"' in xml
    assert 'typeface="DFKai-SB"' in xml


def test_markup_colours_only_the_marked_span(tmp_path: Path) -> None:
    spec = {"deck": {"preset": "zh-classroom"}, "slides": [
        {"layout": "image", "title": "看 **這裡** 和 {positive:那裡}", "notes": "n",
         "source": "s"}]}
    out = tmp_path / "m.pptx"
    build_deck.build(_write(tmp_path, spec), out)
    slide = Presentation(str(out)).slides[0]
    title = next(sh for sh in slide.shapes
                 if sh.has_text_frame and sh.text_frame.text.startswith("看"))
    runs = {r.text: str(r.font.color.rgb) for r in title.text_frame.paragraphs[0].runs}
    assert title.text_frame.text == "看 這裡 和 那裡"
    assert runs["這裡"] == "9C3A24"
    assert runs["那裡"] == "3F6A26"
    assert runs["看 "] == "1F1F1F"


def test_active_section_chip_is_highlighted(tmp_path: Path) -> None:
    spec = {"deck": {"sections": ["甲", "乙"]}, "slides": [
        {"layout": "image", "title": "t", "section": "乙", "notes": "n", "source": "s"}]}
    out = tmp_path / "c.pptx"
    build_deck.build(_write(tmp_path, spec), out)
    slide = Presentation(str(out)).slides[0]
    chip = next(sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text == "乙")
    alert = build_deck.load_theme("default")["tones"]["alert"]["accent"]
    assert str(chip.text_frame.paragraphs[0].runs[0].font.color.rgb) == alert


@pytest.mark.parametrize("bad, message", [
    ({"layout": "nope"}, "unknown layout"),
    ({"layout": "cards", "cards": []}, "1–4 cards"),
    ({"layout": "compare", "panels": [{}]}, "exactly 2 panels"),
    ({"layout": "image", "title": "{purple:x}"}, "unknown tone"),
    ({"layout": "image", "section": "丙"}, "not in deck.sections"),
])
def test_malformed_specs_fail_with_a_named_reason(tmp_path: Path, bad: dict, message: str) -> None:
    spec = {"deck": {"sections": ["甲"]}, "slides": [bad]}
    with pytest.raises(build_deck.SpecError, match=message):
        build_deck.build(_write(tmp_path, spec), tmp_path / "x.pptx")


def test_missing_notes_and_source_are_warned(tmp_path: Path) -> None:
    spec = {"slides": [{"layout": "image", "title": "t"}]}
    warnings = build_deck.build(_write(tmp_path, spec), tmp_path / "w.pptx")
    assert any("no speaker notes" in w for w in warnings)
    assert any("no source footnote" in w for w in warnings)


def test_cli_returns_2_on_bad_spec(tmp_path: Path) -> None:
    path = _write(tmp_path, {"slides": [{"layout": "nope"}]})
    assert build_deck.main([str(path), "-o", str(tmp_path / "x.pptx")]) == 2


def test_generic_demo_builds_with_core_defaults_only(tmp_path: Path) -> None:
    out = tmp_path / "generic.pptx"
    warnings = build_deck.build(GENERIC_DEMO, out)
    prs = Presentation(str(out))
    spec = yaml.safe_load(GENERIC_DEMO.read_text(encoding="utf-8"))
    assert "preset" not in spec["deck"]
    assert {s["layout"] for s in spec["slides"]} == set(build_deck.LAYOUTS)
    assert all("image placeholder" in w for w in warnings), warnings
    xml = prs.slides[1].shapes._spTree.xml
    assert str(prs.slides[1].background.fill.fore_color.rgb) == "FFFFFF"
    assert 'typeface="Calibri"' in xml
    # No preset: nothing course-specific leaks into the deck.
    for course_specific in ("DFKai-SB", "Times New Roman", "［圖］"):
        assert course_specific not in xml
    assert "[Image] an unlabelled fruit" in _texts(prs.slides[1])


def test_preset_supplies_fonts_strings_and_default_theme(tmp_path: Path) -> None:
    spec = {"deck": {"preset": "zh-classroom"}, "slides": [
        {"layout": "image", "title": "t", "image": "missing.png", "image_label": "貓",
         "notes": "n", "source": "s"}]}
    out = tmp_path / "p.pptx"
    build_deck.build(_write(tmp_path, spec), out)
    slide = Presentation(str(out)).slides[0]
    assert str(slide.background.fill.fore_color.rgb) == "FFF8EC"  # warm-paper
    assert "［圖］貓" in _texts(slide)
    assert 'typeface="DFKai-SB"' in slide.shapes._spTree.xml


def test_preset_flag_overrides_the_spec(tmp_path: Path) -> None:
    spec = {"slides": [{"layout": "image", "title": "t", "notes": "n", "source": "s"}]}
    out = tmp_path / "f.pptx"
    build_deck.build(_write(tmp_path, spec), out, theme_override="navy-academic",
                     preset_override="zh-classroom")
    assert str(Presentation(str(out)).slides[0].background.fill.fore_color.rgb) == "FFFFFF"


def test_motif_needs_a_preset_that_defines_it(tmp_path: Path) -> None:
    title = {"layout": "title", "title": "t", "motif": "tree", "notes": "n"}
    with pytest.raises(build_deck.SpecError, match="not provided by preset"):
        build_deck.build(_write(tmp_path, {"slides": [title]}), tmp_path / "x.pptx")
    spec = {"deck": {"preset": "zh-classroom"}, "slides": [title]}
    build_deck.build(_write(tmp_path, spec), tmp_path / "ok.pptx")


def test_course_themes_are_not_core_themes(tmp_path: Path) -> None:
    spec = {"deck": {"theme": "warm-paper"}, "slides": [{"layout": "image", "title": "t"}]}
    with pytest.raises(build_deck.SpecError, match="theme not found"):
        build_deck.build(_write(tmp_path, spec), tmp_path / "x.pptx")


def test_unknown_preset_lists_the_available_ones(tmp_path: Path) -> None:
    spec = {"deck": {"preset": "nope"}, "slides": [{"layout": "image"}]}
    with pytest.raises(build_deck.SpecError, match="available: zh-classroom"):
        build_deck.build(_write(tmp_path, spec), tmp_path / "x.pptx")
