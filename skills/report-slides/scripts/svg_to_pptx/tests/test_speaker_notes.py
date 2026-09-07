"""Tests for carrying speaker notes into the exported deck.

The narrative planner and slide architect both produce speaker notes, and
before this they were dropped silently at export -- a failure with no error to
notice. These tests therefore assert on the *exported deck*, not on the reader
helper alone, so a future change that stops attaching them fails here.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation

from svg_to_pptx.converter import convert_file
from svg_to_pptx.speaker_notes import (
    NOTES_SUFFIX,
    apply_speaker_notes,
    notes_path_for,
    read_speaker_notes,
)

_SVG = (
    '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1200 675">'
    "{desc}"
    '<rect x="10" y="10" width="100" height="50" fill="#123456"/>'
    "</svg>"
)


def _write_slide(directory: Path, name: str, desc: str = "") -> Path:
    """Write one minimal slide SVG, optionally carrying a root `<desc>`."""
    path = directory / name
    path.write_text(_SVG.format(desc=desc), encoding="utf-8")
    return path


def _notes_of(deck: Path) -> list[str]:
    """Return each slide's notes text, empty string when it has none."""
    presentation = Presentation(str(deck))
    texts = []
    for slide in presentation.slides:
        if slide.has_notes_slide:
            texts.append(slide.notes_slide.notes_text_frame.text.strip())
        else:
            texts.append("")
    return texts


def test_notes_path_is_the_svg_stem_plus_suffix(tmp_path: Path) -> None:
    """The sibling name is derived from the slide, not guessed by the caller."""
    assert notes_path_for(tmp_path / "slide-03.svg").name == "slide-03" + NOTES_SUFFIX


def test_sidecar_notes_reach_the_exported_deck(tmp_path: Path) -> None:
    """A `<stem>.notes.md` beside the slide lands in the PPTX notes pane."""
    _write_slide(tmp_path, "slide-01.svg")
    (tmp_path / ("slide-01" + NOTES_SUFFIX)).write_text(
        "Lead with the effect size, then the caveat.", encoding="utf-8"
    )
    deck = tmp_path / "deck.pptx"
    convert_file(str(tmp_path), str(deck))

    assert _notes_of(deck) == ["Lead with the effect size, then the caveat."]


def test_root_desc_is_used_when_no_sidecar_exists(tmp_path: Path) -> None:
    """Notes authored inside the slide file travel with it."""
    _write_slide(tmp_path, "slide-01.svg", desc="<desc>Narration in the SVG.</desc>")
    deck = tmp_path / "deck.pptx"
    convert_file(str(tmp_path), str(deck))

    assert _notes_of(deck) == ["Narration in the SVG."]


def test_sidecar_wins_over_root_desc(tmp_path: Path) -> None:
    """The explicit, separately reviewable source takes precedence."""
    _write_slide(tmp_path, "slide-01.svg", desc="<desc>From the SVG.</desc>")
    (tmp_path / ("slide-01" + NOTES_SUFFIX)).write_text("From the sidecar.", encoding="utf-8")
    deck = tmp_path / "deck.pptx"
    convert_file(str(tmp_path), str(deck))

    assert _notes_of(deck) == ["From the sidecar."]


def test_a_nested_desc_is_not_treated_as_narration(tmp_path: Path) -> None:
    """A `<desc>` inside a shape describes that shape, not the slide."""
    path = tmp_path / "slide-01.svg"
    path.write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1200 675">'
        '<rect x="10" y="10" width="100" height="50" fill="#123456">'
        "<desc>A blue rectangle.</desc></rect></svg>",
        encoding="utf-8",
    )
    deck = tmp_path / "deck.pptx"
    convert_file(str(tmp_path), str(deck))

    assert _notes_of(deck) == [""]


def test_a_slide_without_notes_gets_no_notes_slide(tmp_path: Path) -> None:
    """An empty notes pane reads as deliberate, so none is created."""
    _write_slide(tmp_path, "slide-01.svg")
    deck = tmp_path / "deck.pptx"
    convert_file(str(tmp_path), str(deck))

    assert Presentation(str(deck)).slides[0].has_notes_slide is False


def test_whitespace_only_notes_count_as_absent(tmp_path: Path) -> None:
    """A blank sidecar does not produce an empty notes pane."""
    _write_slide(tmp_path, "slide-01.svg")
    (tmp_path / ("slide-01" + NOTES_SUFFIX)).write_text("   \n\n", encoding="utf-8")

    assert read_speaker_notes(tmp_path / "slide-01.svg") is None


def test_notes_are_matched_per_slide_not_shared(tmp_path: Path) -> None:
    """Each slide keeps its own narration across a multi-slide deck."""
    for index, note in ((1, "First."), (2, None), (3, "Third.")):
        _write_slide(tmp_path, f"slide-{index:02d}.svg")
        if note is not None:
            (tmp_path / f"slide-{index:02d}{NOTES_SUFFIX}").write_text(
                note, encoding="utf-8"
            )
    deck = tmp_path / "deck.pptx"
    convert_file(str(tmp_path), str(deck))

    assert _notes_of(deck) == ["First.", "", "Third."]


def test_multiline_notes_survive_intact(tmp_path: Path) -> None:
    """Paragraph structure is preserved, not flattened to one line."""
    _write_slide(tmp_path, "slide-01.svg")
    (tmp_path / ("slide-01" + NOTES_SUFFIX)).write_text(
        "Open with the question.\n\nThen the method.\nThen the limit.",
        encoding="utf-8",
    )
    deck = tmp_path / "deck.pptx"
    convert_file(str(tmp_path), str(deck))

    assert _notes_of(deck)[0].splitlines()[0] == "Open with the question."
    assert "Then the limit." in _notes_of(deck)[0]


def test_apply_speaker_notes_reports_whether_it_wrote(tmp_path: Path) -> None:
    """The return value is what the export summary counts."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    assert apply_speaker_notes(slide, None) is False
    assert apply_speaker_notes(slide, "Say this.") is True


def test_a_byte_order_mark_does_not_reach_the_notes(tmp_path: Path) -> None:
    """Notes written by a Windows tool start with a BOM; it must not show up.

    Caught by running the real export rather than by reasoning about it:
    PowerShell's `Out-File -Encoding utf8` writes one, and reading as plain
    UTF-8 put a stray U+FEFF at the head of the presenter's first line.
    """
    _write_slide(tmp_path, "slide-01.svg")
    (tmp_path / ("slide-01" + NOTES_SUFFIX)).write_bytes(
        b"\xef\xbb\xbf" + "Open with the effect size.".encode("utf-8")
    )
    deck = tmp_path / "deck.pptx"
    convert_file(str(tmp_path), str(deck))

    notes = _notes_of(deck)[0]
    assert notes == "Open with the effect size."
    assert "\ufeff" not in notes
