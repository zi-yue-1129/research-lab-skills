"""Tests for the exported-deck layout gate.

The finding logic is exercised with layout reports built by hand, so it runs on
the Linux CI image where no PowerPoint exists. A single end-to-end test drives
a real deck through PowerPoint and is skipped elsewhere.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Dict

import pytest

import pptx_com
from validate_pptx_layout import DEFAULT_TOLERANCE_PT, findings_for, main, validate


def _shape(
    name: str,
    height: float,
    bound: float,
    *,
    clipped: bool,
    overflow_x: float = -10.0,
    bounded: bool = True,
) -> Dict[str, Any]:
    """Build one measured shape as `pptx_com.layout` would report it.

    Args:
        name: Shape name.
        height: The shape's height in points.
        bound: The height the laid-out text occupies.
        clipped: Whether the text is cut off rather than growing the shape.
            Expressed through `autosize`, which is what the gate actually
            reads: autosize grows a shape to fit, so its absence is what makes
            vertical overflow a clip.
        overflow_x: Points the text ends past the shape's right edge; negative
            means it finishes inside.
        bounded: Whether the shape draws a fill or outline, so an overflow is
            something a reader can see.
    """
    return {
        "name": name,
        "height": height,
        "width": 200.0,
        "text_bound_height": bound,
        "text_bound_width": 180.0,
        "text_overflow_pt": round(bound - height, 2),
        "text_overflow_x_pt": overflow_x,
        "autosize": 0 if clipped else 1,
        "overflows_box": bound > height or overflow_x > 0,
        "clipped": clipped or overflow_x > 0,
        "has_visible_boundary": bounded,
    }


def _report(*shapes: Dict[str, Any]) -> Dict[str, Any]:
    """Wrap shapes into a one-slide layout report."""
    return {"slide_count": 1, "slides": [{"slide": 1, "shapes": list(shapes)}]}


def test_clipped_text_is_an_error() -> None:
    """Cut-off text loses words, so it fails rather than warns."""
    findings = findings_for(
        _report(_shape("Title", 50.0, 108.0, clipped=True)), DEFAULT_TOLERANCE_PT
    )
    assert len(findings) == 1
    assert findings[0]["rule"] == "pptx-clipped-text"
    assert findings[0]["severity"] == "error"


def test_overflowing_text_is_a_warning() -> None:
    """A growing box stays legible but may collide, so it warns."""
    findings = findings_for(
        _report(_shape("Body", 50.0, 108.0, clipped=False)), DEFAULT_TOLERANCE_PT
    )
    assert findings[0]["rule"] == "pptx-overflowing-text"
    assert findings[0]["severity"] == "warning"
    assert "overlap" in findings[0]["message"]


def test_the_message_carries_measured_and_expected_values() -> None:
    """A finding must be falsifiable without rerunning the gate."""
    message = findings_for(
        _report(_shape("Body", 57.6, 108.0, clipped=True)), DEFAULT_TOLERANCE_PT
    )[0]["message"]
    assert "108.0pt" in message
    assert "57.6pt" in message
    assert "50.4pt over" in message


def test_text_within_its_box_produces_nothing() -> None:
    """A slide that fits is silent."""
    assert findings_for(
        _report(_shape("Body", 108.0, 50.0, clipped=False)), DEFAULT_TOLERANCE_PT
    ) == []


def test_tolerance_suppresses_a_hairline_overflow() -> None:
    """Sub-tolerance overflow is leading, not a defect."""
    report = _report(_shape("Body", 50.0, 50.5, clipped=True))
    assert findings_for(report, 1.0) == []
    assert len(findings_for(report, 0.1)) == 1


def test_shapes_without_text_are_skipped() -> None:
    """A picture or connector has no overflow to measure."""
    assert findings_for(_report({"name": "Picture 1", "height": 10.0}), 1.0) == []


def test_findings_name_their_slide() -> None:
    """A finding has to be actionable, which means saying which slide."""
    report = {
        "slide_count": 2,
        "slides": [
            {"slide": 1, "shapes": [_shape("A", 100.0, 10.0, clipped=False)]},
            {"slide": 2, "shapes": [_shape("B", 10.0, 100.0, clipped=True)]},
        ],
    }
    findings = findings_for(report, DEFAULT_TOLERANCE_PT)
    assert [f["slide"] for f in findings] == [2]
    assert "slide 2" in findings[0]["message"]


def test_main_exits_two_and_reports_blocked_when_unavailable(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    """An unavailable renderer is `blocked`, not a failed deck."""

    def _unavailable(*args: Any, **kwargs: Any) -> Dict[str, Any]:
        raise pptx_com.PowerPointUnavailable("pywin32 is not installed")

    monkeypatch.setattr("validate_pptx_layout.pptx_com.layout", _unavailable)
    assert main(["--pptx", "deck.pptx", "--json"]) == 2

    payload = json.loads(capsys.readouterr().out)
    assert payload["status"] == "blocked"
    assert "pywin32" in payload["blocker"]


def test_main_exits_one_when_findings_were_measured(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    """A clipped deck fails the gate."""
    monkeypatch.setattr(
        "validate_pptx_layout.pptx_com.layout",
        lambda _path: _report(_shape("Title", 50.0, 108.0, clipped=True)),
    )
    assert main(["--pptx", "deck.pptx", "--json"]) == 1
    assert json.loads(capsys.readouterr().out)["status"] == "failed"


def test_main_exits_zero_for_a_clean_deck(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    """A deck that fits passes."""
    monkeypatch.setattr(
        "validate_pptx_layout.pptx_com.layout",
        lambda _path: _report(_shape("Title", 108.0, 50.0, clipped=False)),
    )
    assert main(["--pptx", "deck.pptx", "--json"]) == 0
    assert json.loads(capsys.readouterr().out)["status"] == "passed"


_LIVE = pptx_com.probe()


@pytest.mark.skipif(
    not _LIVE["available"],
    reason=f"PowerPoint COM unavailable: {_LIVE['reason']}",
)
def test_gate_catches_a_real_clip_that_the_svg_linter_cannot(tmp_path: Path) -> None:
    """End to end: a deliberately clipped deck is caught by measuring it.

    This is the case the source-side linter structurally cannot reach --
    `python-pptx` has no font metrics, so nothing before this could know the
    paragraph outgrows its box.
    """
    from pptx import Presentation
    from pptx.enum.text import MSO_AUTO_SIZE
    from pptx.util import Inches, Pt

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(0.8))
    box.text_frame.word_wrap = True
    box.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    box.text_frame.text = (
        "This paragraph is far longer than the box it was placed in, so the "
        "layout engine reflows it well past the declared height."
    )
    for run in box.text_frame.paragraphs[0].runs:
        run.font.size = Pt(18)

    deck = tmp_path / "deck.pptx"
    presentation.save(deck)

    result = validate(deck)
    assert result["status"] == "failed"
    assert result["findings"][0]["rule"] == "pptx-clipped-text"
    assert result["findings"][0]["overflow_pt"] > 0


# --- horizontal overflow --------------------------------------------------


def test_text_escaping_the_right_edge_is_an_error() -> None:
    """Autosize grows height and never width, so this can never come back.

    Measuring only the vertical axis found 3 of 9 real overflows on an actual
    architecture slide; for a diagram, where node boxes have a fixed width and
    their labels outgrow them, horizontal is the dominant failure.
    """
    findings = findings_for(
        _report(_shape("Node", 60.0, 20.0, clipped=False, overflow_x=8.0)),
        DEFAULT_TOLERANCE_PT,
    )
    assert len(findings) == 1
    assert findings[0]["rule"] == "pptx-clipped-text"
    assert findings[0]["severity"] == "error"
    assert findings[0]["axis"] == "horizontal"
    assert "past the shape's right edge" in findings[0]["message"]


def test_a_shape_without_a_visible_boundary_is_not_flagged() -> None:
    """Text past a bare text box's nominal width is how non-wrapping text works.

    Reporting it would bury the real findings under noise nobody can act on --
    there is no box on screen for the reader to see the words leave.
    """
    assert findings_for(
        _report(
            _shape("Caption", 60.0, 20.0, clipped=False, overflow_x=80.0, bounded=False)
        ),
        DEFAULT_TOLERANCE_PT,
    ) == []


def test_both_axes_are_reported_separately() -> None:
    """The fixes differ: shorter text or a wider box, versus fewer lines."""
    findings = findings_for(
        _report(_shape("Node", 20.0, 60.0, clipped=False, overflow_x=8.0)),
        DEFAULT_TOLERANCE_PT,
    )
    assert {f["axis"] for f in findings} == {"horizontal", "vertical"}


def test_horizontal_overflow_respects_the_tolerance() -> None:
    """A hairline is sub-pixel at projection size, not a defect."""
    report = _report(_shape("Node", 60.0, 20.0, clipped=False, overflow_x=0.4))
    assert findings_for(report, DEFAULT_TOLERANCE_PT) == []
