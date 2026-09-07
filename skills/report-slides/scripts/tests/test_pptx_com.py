"""Tests for the PowerPoint COM renderer and layout oracle.

Split deliberately in two. Everything that does not need Office -- shape
measurement, clipped-shape extraction, argument handling, exit codes, report
formatting -- is exercised with fakes so it runs on the Linux CI image, which
is where this file's regressions would otherwise go unnoticed. The handful of
tests that genuinely need a live PowerPoint are skipped everywhere else.

Running this file on Windows prints `Windows fatal exception: code 0x80010108`
(RPC_E_DISCONNECTED) once per PowerPoint shutdown. It is noise, not a failure:
PowerPoint terminates before the automation client's last call returns, and
pytest enables `faulthandler`, which reports the structured exception and
continues. Measured both ways -- calling `Quit()` and merely releasing the
proxy -- it appears identically and leaves no orphan POWERPNT process, so
there is nothing here to fix by changing the teardown.
"""

from __future__ import annotations

import builtins
import json
from pathlib import Path
from typing import Any, Dict, List, Optional

import pytest

from pptx_com import (
    CONVERSION_FORMAT,
    RENDERER_NAME,
    PowerPointUnavailable,
    _emit,
    _measure_shape,
    clipped_shapes,
    layout,
    main,
    probe,
    render,
)


class _FakeTextRange:
    """Stand-in for `TextFrame2.TextRange` with fixed post-layout extents."""

    def __init__(self, text: str, bound_height: float) -> None:
        self.Text = text
        self.BoundLeft = 10.0
        self.BoundTop = 20.0
        self.BoundWidth = 100.0
        self.BoundHeight = bound_height


class _FakeTextFrame:
    """Stand-in for `Shape.TextFrame2`."""

    def __init__(self, text: Optional[str], bound_height: float, autosize: int) -> None:
        self.HasText = text is not None
        self.AutoSize = autosize
        self.TextRange = _FakeTextRange(text or "", bound_height)


class _FakeShape:
    """Duck-typed `Shape`; `_measure_shape` only ever reads these attributes."""

    def __init__(
        self,
        name: str = "TextBox 1",
        height: float = 50.0,
        text: Optional[str] = "hello",
        bound_height: float = 20.0,
        autosize: int = 0,
        has_text_frame: bool = True,
    ) -> None:
        self.Name = name
        self.Left = 1.0
        self.Top = 2.0
        self.Width = 200.0
        self.Height = height
        self.HasTextFrame = has_text_frame
        self.TextFrame2 = _FakeTextFrame(text, bound_height, autosize)


def test_measure_shape_reports_geometry_and_text_extent() -> None:
    """A text shape reports both its box and the laid-out text bounds."""
    measured = _measure_shape(_FakeShape(height=50.0, bound_height=20.0))
    assert measured["name"] == "TextBox 1"
    assert measured["height"] == 50.0
    assert measured["has_text"] is True
    assert measured["text_bound_height"] == 20.0
    assert measured["text_overflow_pt"] == -30.0
    assert measured["clipped"] is False


def test_measure_shape_flags_text_exceeding_a_fixed_box() -> None:
    """Text taller than a non-autosizing shape is a clip."""
    measured = _measure_shape(_FakeShape(height=57.6, bound_height=108.0, autosize=0))
    assert measured["clipped"] is True
    assert measured["text_overflow_pt"] == pytest.approx(50.4)


def test_measure_shape_separates_a_growing_box_from_a_clipped_one() -> None:
    """Under autosize the text stays legible, so it overflows without clipping.

    The shape still grows past its declared bounds and can collide with what
    sits below it, so the overflow is reported rather than discarded.
    """
    measured = _measure_shape(_FakeShape(height=57.6, bound_height=108.0, autosize=1))
    assert measured["clipped"] is False
    assert measured["overflows_box"] is True
    assert measured["text_overflow_pt"] == pytest.approx(50.4)


def test_measure_shape_tolerates_sub_point_overflow() -> None:
    """Rounding-scale overflow is leading, not clipping."""
    measured = _measure_shape(_FakeShape(height=50.0, bound_height=50.4))
    assert measured["clipped"] is False
    assert measured["overflows_box"] is False


def test_measure_shape_handles_a_shape_without_a_text_frame() -> None:
    """A picture or connector reports geometry only."""
    measured = _measure_shape(_FakeShape(has_text_frame=False))
    assert measured["has_text"] is False
    assert "text_bound_height" not in measured


def test_measure_shape_handles_an_empty_text_frame() -> None:
    """An empty placeholder is not measured as text."""
    measured = _measure_shape(_FakeShape(text=None))
    assert measured["has_text"] is False


def test_measure_shape_records_a_measurement_failure_without_raising() -> None:
    """A shape whose text frame errors still yields its geometry."""

    class _Exploding(_FakeShape):
        @property
        def HasTextFrame(self) -> bool:
            raise RuntimeError("COM call failed")

        @HasTextFrame.setter
        def HasTextFrame(self, value: bool) -> None:
            pass

    measured = _measure_shape(_Exploding())
    assert measured["has_text"] is False
    assert "COM call failed" in measured["text_measurement_error"]


def test_clipped_shapes_extracts_only_clipped_entries_with_slide_numbers() -> None:
    """The summary keeps the slide number a finding has to be reported against."""
    report = {
        "slides": [
            {"slide": 1, "shapes": [{"name": "ok", "clipped": False}]},
            {"slide": 2, "shapes": [{"name": "bad", "clipped": True, "text_overflow_pt": 12.0}]},
        ]
    }
    findings = clipped_shapes(report)
    assert len(findings) == 1
    assert findings[0]["slide"] == 2
    assert findings[0]["name"] == "bad"


def test_clipped_shapes_on_an_empty_report() -> None:
    """A report with no slides yields no findings."""
    assert clipped_shapes({}) == []


def test_probe_never_raises_and_explains_unavailability(monkeypatch: pytest.MonkeyPatch) -> None:
    """A non-Windows host is an answer, not an exception."""
    monkeypatch.setattr("pptx_com.sys.platform", "linux")
    report = probe()
    assert report["available"] is False
    assert report["name"] == RENDERER_NAME
    assert "Windows" in report["reason"]


def test_probe_reports_missing_pywin32(monkeypatch: pytest.MonkeyPatch) -> None:
    """A Windows host without pywin32 names the missing package."""
    monkeypatch.setattr("pptx_com.sys.platform", "win32")
    real_import = builtins.__import__

    def _fail(name: str, *args: Any, **kwargs: Any) -> Any:
        if name.startswith("win32com"):
            raise ImportError("no pywin32")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", _fail)
    report = probe()
    assert report["available"] is False
    assert "pywin32" in report["reason"]


def test_main_requires_a_deck_for_render(capsys: pytest.CaptureFixture[str]) -> None:
    """`--render` without `--pptx` is a usage error, not a render attempt."""
    assert main(["--render", "--out", "x"]) == 1
    assert "--pptx is required" in capsys.readouterr().err


def test_main_requires_an_output_directory_for_render(capsys: pytest.CaptureFixture[str]) -> None:
    """`--render` without `--out` is a usage error."""
    assert main(["--render", "--pptx", "deck.pptx"]) == 1
    assert "--out is required" in capsys.readouterr().err


def test_main_requires_a_deck_for_layout(capsys: pytest.CaptureFixture[str]) -> None:
    """`--layout` without `--pptx` is a usage error."""
    assert main(["--layout"]) == 1
    assert "--pptx is required" in capsys.readouterr().err


def test_main_returns_two_when_the_host_cannot_render(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    """An unavailable renderer exits 2 so a caller can fall back to LibreOffice."""

    def _unavailable(*args: Any, **kwargs: Any) -> Dict[str, Any]:
        raise PowerPointUnavailable("pywin32 is not installed")

    monkeypatch.setattr("pptx_com.layout", _unavailable)
    assert main(["--layout", "--pptx", "deck.pptx", "--json"]) == 2
    payload = json.loads(capsys.readouterr().out)
    assert payload["error"] == "PowerPointUnavailable"
    assert "pywin32" in payload["message"]


def test_main_probe_emits_json(capsys: pytest.CaptureFixture[str]) -> None:
    """`--probe --json` is machine-readable on every platform."""
    assert main(["--probe", "--json"]) == 0
    payload = json.loads(capsys.readouterr().out)
    assert payload["name"] == RENDERER_NAME
    assert isinstance(payload["available"], bool)


def test_emit_summarises_a_render_report(capsys: pytest.CaptureFixture[str]) -> None:
    """The human summary lists every rendered PNG."""
    _emit(
        {
            "slide_count": 2,
            "renderer": {"name": RENDERER_NAME},
            "rendered_png_paths": ["a/slide-01.png", "a/slide-02.png"],
        },
        as_json=False,
    )
    out = capsys.readouterr().out
    assert "rendered 2 slide(s)" in out
    assert "slide-02.png" in out


def test_emit_summarises_clipped_shapes(capsys: pytest.CaptureFixture[str]) -> None:
    """The layout summary leads with the clip count."""
    _emit(
        {
            "slide_count": 1,
            "slides": [
                {"slide": 1, "shapes": [{"name": "TextBox 1", "clipped": True, "text_overflow_pt": 50.4}]}
            ],
        },
        as_json=False,
    )
    out = capsys.readouterr().out
    assert "clipped shapes: 1" in out
    assert "50.4pt" in out


# --- live-PowerPoint integration ------------------------------------------
# Probed once: evaluating this per test would start and quit PowerPoint
# repeatedly during collection.
_LIVE = probe()

requires_powerpoint = pytest.mark.skipif(
    not _LIVE["available"],
    reason=f"PowerPoint COM unavailable: {_LIVE['reason']}",
)


@pytest.fixture
def sample_deck(tmp_path: Path) -> Path:
    """Build a two-slide deck whose second slide deliberately overflows."""
    from pptx import Presentation
    from pptx.enum.text import MSO_AUTO_SIZE
    from pptx.util import Inches, Pt

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    first = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = first.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11), Inches(1.5))
    box.text_frame.text = "Fits comfortably"
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(44)

    second = presentation.slides.add_slide(presentation.slide_layouts[6])
    tight = second.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(0.8))
    tight.text_frame.word_wrap = True
    # python-pptx puts `spAutoFit` on a new textbox, which makes PowerPoint
    # grow the shape instead of cutting the text off. Pin the size so this is
    # the genuine clip the oracle exists to catch.
    tight.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    tight.text_frame.text = (
        "This paragraph is far longer than the box it was placed in, so the "
        "layout engine reflows it well past the declared height."
    )
    for run in tight.text_frame.paragraphs[0].runs:
        run.font.size = Pt(18)

    path = tmp_path / "deck.pptx"
    presentation.save(path)
    return path


@requires_powerpoint
def test_render_writes_one_png_per_slide(sample_deck: Path, tmp_path: Path) -> None:
    """Rendering produces exactly one non-empty PNG per slide, plus a PDF."""
    out = tmp_path / "renders"
    report = render(sample_deck, out, width=960, height=540)

    assert report["slide_count"] == 2
    assert report["renderer"]["name"] == RENDERER_NAME
    assert report["renderer"]["conversion_format"] == CONVERSION_FORMAT
    assert report["renderer"]["version"]

    pngs: List[Path] = [Path(p) for p in report["rendered_png_paths"]]
    assert [p.name for p in pngs] == ["slide-01.png", "slide-02.png"]
    assert all(p.is_file() and p.stat().st_size > 0 for p in pngs)
    assert report["conversion_artifacts"]
    assert Path(report["conversion_artifacts"][0]).is_file()


@requires_powerpoint
def test_render_without_pdf_falls_back_to_the_png_list(sample_deck: Path, tmp_path: Path) -> None:
    """`conversion_artifacts` stays non-empty, as the review record requires."""
    report = render(sample_deck, tmp_path / "renders", width=480, height=270, pdf=False)
    assert report["conversion_artifacts"] == report["rendered_png_paths"]


@requires_powerpoint
def test_layout_detects_the_deliberate_overflow(sample_deck: Path) -> None:
    """The oracle finds the clip that python-pptx structurally cannot see."""
    report = layout(sample_deck)
    assert report["slide_count"] == 2

    findings = clipped_shapes(report)
    assert [f["slide"] for f in findings] == [2]
    assert findings[0]["text_overflow_pt"] > 0


@requires_powerpoint
def test_render_rejects_a_missing_deck(tmp_path: Path) -> None:
    """A missing deck is a typed, reportable blocker."""
    with pytest.raises(PowerPointUnavailable, match="deck not found"):
        render(tmp_path / "absent.pptx", tmp_path / "out")


class _FakePageSetup:
    """Page setup reporting a slide's dimensions in points."""

    def __init__(self, width: float, height: float) -> None:
        self.SlideWidth = width
        self.SlideHeight = height


class _FakePresentation:
    """Just enough presentation for `_export_size` to read a page shape."""

    def __init__(self, width: float, height: float) -> None:
        self.PageSetup = _FakePageSetup(width, height)


def test_export_size_derives_height_from_a_widescreen_deck() -> None:
    """A 16:9 deck renders at the default long edge, undistorted."""
    from pptx_com import DEFAULT_LONG_EDGE, _export_size

    assert _export_size(_FakePresentation(960.0, 540.0), None, None) == (
        DEFAULT_LONG_EDGE,
        DEFAULT_LONG_EDGE * 9 // 16,
    )


def test_export_size_does_not_stretch_a_four_three_deck() -> None:
    """The defect this replaced: a 4:3 deck forced into a 16:9 frame.

    The visual gate inspects these pixels to judge what the reader will see, so
    handing it a horizontally stretched picture undermines the whole gate.
    """
    from pptx_com import _export_size

    width, height = _export_size(_FakePresentation(720.0, 540.0), None, None)
    assert round(width / height, 3) == round(720.0 / 540.0, 3)


def test_export_size_handles_a_portrait_deck() -> None:
    """A taller-than-wide deck puts the long edge on the height."""
    from pptx_com import DEFAULT_LONG_EDGE, _export_size

    width, height = _export_size(_FakePresentation(540.0, 720.0), None, None)
    assert height == DEFAULT_LONG_EDGE
    assert round(width / height, 3) == round(540.0 / 720.0, 3)


def test_export_size_derives_the_missing_side_from_one_given() -> None:
    """Capping one dimension must not skew the other."""
    from pptx_com import _export_size

    assert _export_size(_FakePresentation(960.0, 540.0), 1280, None) == (1280, 720)
    assert _export_size(_FakePresentation(960.0, 540.0), None, 720) == (1280, 720)


def test_export_size_honours_both_dimensions_when_given() -> None:
    """An explicit size is obeyed verbatim, distortion included.

    A caller comparing against a fixed-size reference sometimes needs exactly
    that, so the override is not second-guessed.
    """
    from pptx_com import _export_size

    assert _export_size(_FakePresentation(720.0, 540.0), 1920, 1080) == (1920, 1080)
